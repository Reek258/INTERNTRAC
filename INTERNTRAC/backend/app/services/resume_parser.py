import os
from pypdf import PdfReader
from docx import Document

class ResumeParser:
    @staticmethod
    def extract_text_from_file(file_path: str) -> str:
        """
        Extract plain text from PDF or DOCX resume.
        Raises ValueError if format is unsupported or file is corrupted/unreadable.
        """
        if not os.path.exists(file_path):
            raise ValueError("Resume file not found.")

        ext = os.path.splitext(file_path)[1].lower()
        if ext not in [".pdf", ".docx"]:
            raise ValueError("Only PDF and DOCX resume files are supported.")

        text = ""
        try:
            if ext == ".pdf":
                reader = PdfReader(file_path)
                if len(reader.pages) == 0:
                    raise ValueError("The uploaded PDF has no readable pages.")
                for page in reader.pages:
                    extracted = page.extract_text()
                    if extracted:
                        text += extracted + "\n"
            elif ext == ".docx":
                doc = Document(file_path)
                paragraphs_text = [p.text for p in doc.paragraphs if p.text]
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            if cell.text:
                                paragraphs_text.append(cell.text)
                text = "\n".join(paragraphs_text)

            text = text.strip()
            if not text:
                raise ValueError("We could not read any text from this resume. Please upload a valid PDF or DOCX file.")
            return text
        except Exception as e:
            if isinstance(e, ValueError):
                raise e
            raise ValueError(f"We could not read this resume. Please upload a valid PDF or DOCX file. (Error: {str(e)})")

    @staticmethod
    def extract_skills_from_text(text: str) -> list[str]:
        """
        Extract common technical skills matching keywords in resume.
        """
        text_lower = text.lower()
        common_skills = [
            "Python", "JavaScript", "TypeScript", "React", "React.js", "Node.js", "Express",
            "FastAPI", "Django", "Flask", "SQL", "PostgreSQL", "MySQL", "SQLite", "MongoDB",
            "Redis", "Docker", "Kubernetes", "AWS", "Azure", "GCP", "Git", "GitHub", "Linux",
            "HTML", "CSS", "Tailwind CSS", "Bootstrap", "Next.js", "Vue.js", "Angular",
            "C", "C++", "Java", "Spring Boot", "Go", "Rust", "GraphQL", "REST API", "CI/CD",
            "Machine Learning", "Deep Learning", "NLP", "Pandas", "NumPy", "Scikit-Learn",
            "TensorFlow", "PyTorch", "Data Structures", "Algorithms", "Figma", "UI/UX"
        ]

        found = []
        for skill in common_skills:
            skill_lower = skill.lower()
            if skill_lower in text_lower:
                found.append(skill)
        return found

    @classmethod
    def extract_text(cls, file_path: str) -> str:
        return cls.extract_text_from_file(file_path)

resume_parser = ResumeParser()
